#!/usr/bin/env python3
"""Transform the Zenodo IoTS_Dataset into BIM v1 packages.

Dataset: https://zenodo.org/records/10440967 (Tang et al., Heliyon 2024)
Paper: 'IoT service composition based on improved Shuffled Frog Leaping Algorithm',
       Heliyon, 2024. DOI: 10.1016/j.heliyon.2024.e28087.
Archive contains 6 benchmark scales:
  IoTS10X50, IoTS10X100, IoTS20X50, IoTS20X100, IoTS30X50, IoTS30X100
"""
from __future__ import annotations

import argparse
import io
import json
import re
import sys
import zipfile
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "openbinding-gateway/src"))

from openbinding_gateway.v1.compiler import compile_instance
from openbinding_gateway.v1.package import InstancePackage


def read_task_workbook(wb_bytes: bytes) -> list[dict[str, float]]:
    wb = load_workbook(io.BytesIO(wb_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []

    header = [str(x).strip().lower() if x is not None else "" for x in rows[0]]
    col_map = {}
    for idx, h in enumerate(header):
        if "time" in h:
            col_map["execution_time"] = idx
        elif "cost" in h or "price" in h:
            col_map["cost"] = idx
        elif "credibility" in h or "reputation" in h or "rep" in h:
            col_map["reputation"] = idx
        elif "reliability" in h or "rel" in h:
            col_map["reliability"] = idx

    if len(col_map) < 4:
        col_map = {"execution_time": 1, "cost": 2, "reputation": 3, "reliability": 4}

    candidates = []
    for r in rows[1:]:
        try:
            time_val = float(r[col_map["execution_time"]])
            cost_val = float(r[col_map["cost"]])
            rep_val = float(r[col_map["reputation"]])
            rel_val = float(r[col_map["reliability"]])
            candidates.append({
                "execution_time": time_val,
                "cost": cost_val,
                "reputation": rep_val,
                "reliability": rel_val,
            })
        except (TypeError, ValueError, IndexError):
            continue
    return candidates


def transform_scale(scale_name: str, tasks_data: dict[str, list[dict[str, float]]]) -> InstancePackage:
    clean_name = f"iots_{scale_name.lower().replace('iots', '')}"
    task_keys = sorted(tasks_data.keys(), key=lambda t: int(re.search(r"\d+", t).group(0)))

    all_cands = [c for c_list in tasks_data.values() for c in c_list]

    feature_keys = ["execution_time", "cost", "reputation", "reliability"]
    minima = {k: min(c[k] for c in all_cands) for k in feature_keys}
    maxima = {k: max(c[k] for c in all_cands) for k in feature_keys}

    feature_defs = {
        "execution_time": {
            "direction": "minimize",
            "scope": "invocation",
            "aggregation": "sum",
            "domain": {"kind": "real", "minimum": 0.0, "maximum": max(float(maxima["execution_time"]) * 1.5, 1000.0)},
        },
        "cost": {
            "direction": "minimize",
            "scope": "selectedCandidate",
            "aggregation": "sum",
            "domain": {"kind": "real", "minimum": 0.0, "maximum": max(float(maxima["cost"]) * 1.5, 1000.0)},
        },
        "reputation": {
            "direction": "maximize",
            "scope": "selectedCandidate",
            "aggregation": "sum",
            "domain": {"kind": "real", "minimum": 0.0, "maximum": max(float(maxima["reputation"]) * 1.5, 100.0)},
        },
        "reliability": {
            "direction": "maximize",
            "scope": "selectedCandidate",
            "aggregation": "product",
            "domain": {"kind": "real", "minimum": 0.0, "maximum": float(max(1.0, maxima["reliability"] + 0.1))},
        },
    }

    tasks_dict = {tid: f"service/{tid}" for tid in task_keys}
    candidates_dict = {}
    providers_dict = {"iots_source_provider": {}}

    for tid in task_keys:
        cands = tasks_data[tid]
        for c_idx, c_feat in enumerate(cands):
            c_id = f"c_{tid}_{c_idx + 1}"
            candidates_dict[c_id] = {
                "provides": f"service/{tid}",
                "provider": {"resource": "catalog", "id": "iots_source_provider"},
                "features": c_feat,
            }

    workflow_seq = [{"task": {"resource": "application", "id": tid}} for tid in task_keys]

    inst = {
        "apiVersion": "bim/v1",
        "kind": "Instance",
        "metadata": {
            "name": clean_name,
            "version": "1.0.0",
            "description": f"IoTS benchmark instance for scale {scale_name} ({len(task_keys)} tasks)",
            "annotations": {
                "source": "https://zenodo.org/records/10440967",
                "paper": {
                    "authors": "M. Tang et al.",
                    "title": "IoT service composition based on improved Shuffled Frog Leaping Algorithm",
                    "journal": "Heliyon",
                    "year": 2024,
                    "doi": "10.1016/j.heliyon.2024.e28087",
                },
                "license": "Creative Commons Attribution 4.0 International",
                "tasks_count": len(task_keys),
                "candidates_per_task": len(tasks_data[task_keys[0]]) if task_keys else 0,
            },
        },
        "spec": {
            "profile": "qos-binding/v1",
            "resources": {
                "application": {"application": "application.json"},
                "candidateCatalog": {"catalog": "candidates.json"},
                "constraintSet": {"constraints": "constraints.json"},
                "optimization": {"optimization": "optimization.json"},
            },
        },
    }

    app = {
        "apiVersion": "qos-binding/v1",
        "kind": "Application",
        "metadata": {"name": f"{clean_name}_application"},
        "spec": {
            "tasks": tasks_dict,
            "features": feature_defs,
            "workflow": {"sequence": workflow_seq},
        },
    }

    cat = {
        "apiVersion": "qos-binding/v1",
        "kind": "CandidateCatalog",
        "metadata": {"name": f"{clean_name}_catalog"},
        "spec": {
            "providers": providers_dict,
            "featureBindings": {k: {"resource": "application", "id": k} for k in feature_defs},
            "candidates": candidates_dict,
        },
    }

    cons = {
        "apiVersion": "qos-binding/v1",
        "kind": "ConstraintSet",
        "metadata": {"name": f"{clean_name}_constraints"},
        "spec": {"constraints": {}},
    }

    opt = {
        "apiVersion": "qos-binding/v1",
        "kind": "Optimization",
        "metadata": {"name": f"{clean_name}_optimization"},
        "spec": {
            "criteria": [
                {
                    "id": k,
                    "feature": {"resource": "application", "id": k},
                    "direction": feature_defs[k]["direction"],
                    "normalize": {"min": 0.0, "max": float(maxima[k]), "clamp": True},
                }
                for k in feature_defs
            ]
        },
    }

    files = {
        "instance.json": (json.dumps(inst, indent=2, sort_keys=True) + "\n").encode(),
        "application.json": (json.dumps(app, indent=2, sort_keys=True) + "\n").encode(),
        "candidates.json": (json.dumps(cat, indent=2, sort_keys=True) + "\n").encode(),
        "constraints.json": (json.dumps(cons, indent=2, sort_keys=True) + "\n").encode(),
        "optimization.json": (json.dumps(opt, indent=2, sort_keys=True) + "\n").encode(),
    }
    return InstancePackage(files)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source",
        type=Path,
        default=ROOT / "datasets/sources/iots",
        help="Path to uncompressed iots directory or zip file",
    )
    parser.add_argument(
        "--target-dir",
        type=Path,
        default=ROOT / "datasets/05_iots",
        help="Destination directory for transformed packages",
    )
    args = parser.parse_args()

    if not args.source.exists():
        raise SystemExit(f"Source not found: {args.source}")

    args.target_dir.mkdir(parents=True, exist_ok=True)

    scales: dict[str, dict[str, bytes]] = {}
    if args.source.is_dir():
        for xp in sorted(args.source.glob("**/*.xlsx")):
            scale_dir = xp.parent.name
            t_name = re.sub(r"[^A-Za-z0-9_]+", "", xp.stem).lower()
            if scale_dir not in scales:
                scales[scale_dir] = {}
            scales[scale_dir][t_name] = xp.read_bytes()
    else:
        zf = zipfile.ZipFile(args.source)
        for fn in zf.namelist():
            if fn.endswith(".xlsx") and "/" in fn:
                parts = fn.split("/")
                scale_dir = parts[1]
                task_fn = parts[2]
                t_name = re.sub(r"[^A-Za-z0-9_]+", "", Path(task_fn).stem).lower()
                if scale_dir not in scales:
                    scales[scale_dir] = {}
                scales[scale_dir][t_name] = zf.read(fn)

    print(f"Found {len(scales)} IoTS scale suites: {list(scales.keys())}")
    compiled_count = 0

    for scale_name, task_files in sorted(scales.items()):
        tasks_data = {}
        for t_name, b in task_files.items():
            cands = read_task_workbook(b)
            if cands:
                tasks_data[t_name] = cands

        pkg = transform_scale(scale_name, tasks_data)
        compiled = compile_instance(pkg)

        clean_name = f"iots_{scale_name.lower().replace('iots', '')}"
        out_dir = args.target_dir / clean_name
        out_dir.mkdir(parents=True, exist_ok=True)
        for rel_path, content in pkg.files.items():
            (out_dir / rel_path).write_bytes(content)

        compiled_count += 1
        print(f"  [OK] {clean_name}: {len(tasks_data)} tasks, compiled digest {compiled.digest}")

    print(f"Successfully transformed and compiled {compiled_count}/{len(scales)} IoTS scale suites!")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
